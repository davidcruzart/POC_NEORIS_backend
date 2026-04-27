from io import BytesIO

from docx import Document

from app.interfaces.exporter import Exporter


class DocxExporter(Exporter):
    def export(self, text: str) -> bytes:
        document = Document()
        document.add_heading("Resumen generado", level=1)
        document.add_paragraph(text)

        buffer = BytesIO()
        document.save(buffer)
        buffer.seek(0)

        return buffer.getvalue()